import pytest
from pathlib import Path

from backend.core.tools.document_tools import document_create, document_open
from tests.test_document_candidate_service import document_native_runtime, document_native_launcher, source_runtime, isolated_python


def test_office_and_pdf_adapters_use_one_bounded_contract(tmp_path_factory: Path):
    workspace = tmp_path_factory

    from docx import Document
    word = Document()
    word.add_paragraph("Word adapter text")
    word.save(workspace / "sample.docx")

    from pptx import Presentation
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "PowerPoint adapter text"
    deck.save(workspace / "sample.pptx")

    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = "Evidence"
    sheet.append(["Excel", "adapter", 42])
    book.save(workspace / "sample.xlsx")

    from pypdf import PdfWriter
    pdf = PdfWriter()
    pdf.add_blank_page(width=100, height=100)
    with (workspace / "sample.pdf").open("wb") as stream:
        pdf.write(stream)

    word_result = document_open({"_workspace": str(workspace), "path": "sample.docx"})
    slides_result = document_open({"_workspace": str(workspace), "path": "sample.pptx"})
    sheet_result = document_open({"_workspace": str(workspace), "path": "sample.xlsx"})
    pdf_result = document_open({"_workspace": str(workspace), "path": "sample.pdf"})

    assert "Word adapter text" in word_result["content"]
    assert "PowerPoint adapter text" in slides_result["content"]
    assert "Excel\tadapter\t42" in sheet_result["content"]
    assert pdf_result["format"] == "pdf" and "error" not in pdf_result
    assert all("next_offset" in item for item in (word_result, slides_result, sheet_result, pdf_result))


def test_unsupported_document_returns_recoverable_catalog_error(tmp_path_factory: Path):
    workspace = tmp_path_factory
    (workspace / "sample.bin").write_bytes(b"fixture")

    result = document_open({"_workspace": str(workspace), "path": "sample.bin"})

    assert result["error"] == "DOCUMENT_FORMAT_UNSUPPORTED"
    assert result["error_info"]["catalog_id"] == "GITGO-E3401"
    assert "supported" in result["error_info"]["details"]


def test_document_create_round_trips_common_output_formats(tmp_path_factory: Path, document_native_launcher):
    workspace = tmp_path_factory
    cases = {
        "note.md": "# Heading\nDurable markdown",
        "brief.docx": "Durable Word text",
        "deck.pptx": "Durable slide text",
        "table.xlsx": "name\tvalue\nGitgo\t42",
        "report.pdf": "Durable PDF text",
    }

    for path, content in cases.items():
        from tests.test_file_feedback_contract import context
        from backend.core.tools.catalog import build_workspace_tools
        from backend.core.loop.tool_pipeline import ToolPipeline
        ctx = context(workspace, ["document_create"])
        tool = build_workspace_tools(workspace)["document_create"]
        result = ToolPipeline().execute({"name": "document_create", "args": {
            "path": path, "title": "Gitgo output", "content": content}}, tool, ctx, path)
        created = result.data
        assert not result.is_error, result.formatted
        if not path.endswith(".md"):
            assert result.receipt["candidate_render"]["launch_state"] == "started"
            assert result.receipt["candidate_render"]["cleanup"]["complete"] is True
            assert result.receipt["candidate_render"]["access_retired"] is True
            assert result.receipt["runtime_cleanup"]["complete"] is True
            assert result.receipt["effect_state"] == "committed"
        assert created.get("action") == "created", (path, created)
        assert created["sha256"]
        assert (workspace / path).is_file()
        opened = document_open({"_workspace": str(workspace), "path": path})
        assert "error" not in opened
        expected = "Heading" if path.endswith(".md") else content.split("\n", 1)[0].split("\t", 1)[0]
        assert expected in opened["content"]


def test_document_create_requires_hash_bound_overwrite(tmp_path_factory: Path):
    workspace = tmp_path_factory
    target = workspace / "existing.md"
    target.write_text("original", encoding="utf-8")

    rejected = document_create({
        "_workspace": str(workspace), "path": "existing.md", "content": "changed",
    })

    assert rejected["error"] == "DOCUMENT_WRITE_FAILED"
    assert rejected["error_info"]["catalog_id"] == "GITGO-E3405"
    assert target.read_text(encoding="utf-8") == "original"


def test_doc_converter_rejects_workspace_and_relative_path_candidates(tmp_path_factory, monkeypatch):
    import os
    from backend.core.tools.document_limits import verified_antiword
    workspace = tmp_path_factory / 'workspace'
    workspace.mkdir()
    name = 'antiword.exe' if os.name == 'nt' else 'antiword'
    fake = workspace / name
    fake.write_bytes(b'workspace executable')
    fake.chmod(0o755)
    monkeypatch.chdir(workspace)
    monkeypatch.setenv('PATH', str(workspace) + os.pathsep + '.')
    import hashlib
    monkeypatch.setenv("GITGO_ANTIWORD_SHA256", hashlib.sha256(fake.read_bytes()).hexdigest())
    with pytest.raises(ValueError, match="CONVERTER_IDENTITY_UNVERIFIED"):
        verified_antiword(workspace)
    (workspace / 'legacy.doc').write_bytes(b'legacy document')
    result = document_open({'_workspace': str(workspace), 'path': 'legacy.doc'})
    assert result['error'] == 'DOCUMENT_FORMAT_UNSUPPORTED'


def test_doc_converter_selects_absolute_host_installation(tmp_path_factory, monkeypatch):
    import os
    from backend.core.tools.document_limits import verified_antiword
    workspace = tmp_path_factory / 'workspace'
    runtime = tmp_path_factory / 'runtime'
    workspace.mkdir(); runtime.mkdir()
    converter = runtime / ('antiword.exe' if os.name == 'nt' else 'antiword')
    converter.write_bytes(b'installed executable identity'); converter.chmod(0o755)
    monkeypatch.setenv('PATH', str(runtime))
    import hashlib
    digest = hashlib.sha256(converter.read_bytes()).hexdigest()
    monkeypatch.setenv("GITGO_ANTIWORD_SHA256", digest)
    assert verified_antiword(workspace) == (converter.resolve(), digest)


def test_doc_converter_receives_private_eof(tmp_path_factory, monkeypatch):
    import subprocess
    from types import SimpleNamespace
    from backend.core.tools import document_tools, document_limits
    workspace = tmp_path_factory / 'workspace'
    workspace.mkdir()
    path = workspace / 'legacy.doc'
    path.write_bytes(b'legacy document')
    converter = tmp_path_factory / 'antiword'
    converter.write_bytes(b'fixture converter')
    monkeypatch.setattr(document_limits, 'verified_antiword', lambda _: (converter, 'fixture'))
    monkeypatch.setattr(document_limits, 'check_converter', lambda *args: None)
    def convert(argv, **kwargs):
        assert kwargs['stdin'] == subprocess.DEVNULL
        assert argv == [str(converter), str(path)]
        return SimpleNamespace(returncode=0)
    monkeypatch.setattr(document_tools.subprocess, 'Popen', convert)
    monkeypatch.setattr('backend.core.process_control.attach_kill_job', lambda _: None)
    monkeypatch.setattr('backend.core.process_output.collect_output',
        lambda *args, **kwargs: ('converted document', '', False))
    assert document_open({'_workspace': str(workspace), 'path': path.name})['content'] == 'converted document'
